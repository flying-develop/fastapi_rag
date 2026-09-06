"""Extract plain text from uploaded files — PDF/DOCX/XLSX only.

Deliberately narrow scope: only the three document formats the "Работа с
файлами" milestone asks for, and only plain text — no formatting/structure
is preserved (no tables-as-tables, no styles). Everything else (images,
audio/video, or any other media) is explicitly ignored rather than
attempted — see `parse_to_text()`.
"""

import io
import logging

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

logger = logging.getLogger(__name__)

_PDF_CONTENT_TYPE = "application/pdf"
_DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_xlsx_text(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(lines)


_EXTRACTORS = {
    _PDF_CONTENT_TYPE: _extract_pdf_text,
    _DOCX_CONTENT_TYPE: _extract_docx_text,
    _XLSX_CONTENT_TYPE: _extract_xlsx_text,
}


def parse_to_text(content_type: str, data: bytes) -> tuple[str | None, str]:
    """Return `(extracted_text, parse_status)` for an uploaded file.

    `parse_status` is one of:
    - `"success"` — `content_type` is PDF/DOCX/XLSX and text was extracted.
    - `"skipped"` — `content_type` is anything else (images, audio/video,
      or any other media). No extraction is attempted for these — this is
      a deliberate milestone requirement, not a missing feature.
    - `"failed"` — `content_type` is a supported one, but the file's bytes
      couldn't actually be parsed (corrupted/invalid file). Never raises:
      the caller (`FileService.upload_file`) must still save the file
      even when its content can't be parsed.

    Synchronous by design — `pypdf`/`python-docx`/`openpyxl` are all
    sync libraries; the caller is expected to run this off the event
    loop (`asyncio.to_thread`), the same way sync tool functions run in
    a threadpool for `invoke_with_tools()` (see `docs/tool-calling.md`).
    """
    extractor = _EXTRACTORS.get(content_type)
    if extractor is None:
        logger.info(
            "parsing skipped — unsupported content type",
            extra={"content_type": content_type},
        )
        return None, "skipped"

    try:
        text = extractor(data)
    except Exception as exc:
        logger.warning(
            "parsing failed",
            extra={"content_type": content_type, "error_type": type(exc).__name__},
        )
        return None, "failed"

    logger.info(
        "parsing succeeded",
        extra={"content_type": content_type, "extracted_length": len(text)},
    )
    return text, "success"
