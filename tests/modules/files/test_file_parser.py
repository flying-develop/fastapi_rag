"""Tests for `parse_to_text()` — PDF/DOCX/XLSX generated in-memory
(`reportlab`/`python-docx`/`openpyxl`, no static binary fixtures in git),
plus the "unsupported type" and "corrupted file" edge cases."""

import io

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.modules.files.services.file_parser import parse_to_text


def _make_pdf_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(100, 700, text)
    pdf.save()
    return buffer.getvalue()


def _make_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx_bytes(value: str) -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_to_text_extracts_pdf_text() -> None:
    data = _make_pdf_bytes("Hello PDF")

    text, status = parse_to_text("application/pdf", data)

    assert status == "success"
    assert text is not None
    assert "Hello PDF" in text


def test_parse_to_text_extracts_docx_text() -> None:
    data = _make_docx_bytes("Hello DOCX")

    text, status = parse_to_text(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data,
    )

    assert status == "success"
    assert text is not None
    assert "Hello DOCX" in text


def test_parse_to_text_extracts_xlsx_text() -> None:
    data = _make_xlsx_bytes("Hello XLSX")

    text, status = parse_to_text(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", data
    )

    assert status == "success"
    assert text is not None
    assert "Hello XLSX" in text


def test_parse_to_text_skips_images_and_other_media() -> None:
    text, status = parse_to_text("image/png", b"\x89PNG fake bytes, never decoded")

    assert status == "skipped"
    assert text is None


def test_parse_to_text_returns_failed_for_corrupted_supported_type() -> None:
    text, status = parse_to_text("application/pdf", b"not a real pdf")

    assert status == "failed"
    assert text is None
